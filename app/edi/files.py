"""v0.8: EDI puede leer, resumir y corregir archivos — pero SÓLO dentro de una carpeta que tú autorizaste
(`config.FILES_DIR`, normalmente montada en docker-compose a una carpeta real de tu Mac, p. ej.
~/Documents/EDI-files). Nunca toca el resto del disco. Las ediciones nunca pisan el original: se guardan
con sufijo "_EDI" junto al archivo, para que siempre puedas comparar y decidir."""
import difflib
import logging
import re
from pathlib import Path

from . import config, db, providers

log = logging.getLogger("edi.files")

TEXT_EXT = {".txt", ".md", ".py", ".js", ".ts", ".json", ".yaml", ".yml", ".html", ".css", ".sh", ".csv", ".java",
            ".c", ".cpp", ".go", ".rb", ".php", ".sql"}
CODE_EXT = {".py", ".js", ".ts", ".java", ".c", ".cpp", ".go", ".rb", ".php", ".sh", ".sql"}


class NotAllowed(Exception):
    pass


def _safe(rel: str) -> Path:
    """Resuelve una ruta DENTRO de la carpeta autorizada; nunca deja salir con '../' ni rutas absolutas."""
    base = config.FILES_DIR.resolve()
    p = (base / rel.lstrip("/\\")).resolve()
    if base not in p.parents and p != base:
        raise NotAllowed(f"{rel} está fuera de la carpeta autorizada de EDI")
    return p


def list_files(limit: int = 200) -> list[dict]:
    base = config.FILES_DIR
    out = []
    for p in sorted(base.rglob("*")):
        if p.is_file() and not p.name.startswith("."):
            out.append({"name": p.name, "rel": str(p.relative_to(base)), "size": p.stat().st_size,
                        "ext": p.suffix.lower()})
    return out[:limit]


def find(name: str) -> list[dict]:
    """Busca por nombre aproximado dentro de la carpeta autorizada ('el pdf de la actividad', 'reporte.docx')."""
    q = re.sub(r"[^\w\s.]", " ", name.lower()).strip()
    q_words = [w for w in q.split() if len(w) > 1]
    scored = []
    for f in list_files(1000):
        n = f["name"].lower()
        if q == n or q == Path(n).stem:
            sc = 5
        elif all(w in n for w in q_words) and q_words:
            sc = 3
        else:
            sc = 2 if difflib.SequenceMatcher(None, q, n).ratio() >= 0.6 else 0
        if sc:
            scored.append((sc, f))
    scored.sort(key=lambda x: -x[0])
    return [f for _, f in scored[:5]]


def read_text(rel: str, limit: int = 20000) -> str:
    return _safe(rel).read_text(encoding="utf-8", errors="ignore")[:limit]


def read_pdf(rel: str, limit: int = 20000) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("Falta la librería pypdf (vuelve a construir la imagen: ./scripts/edi.sh rebuild)")
    reader = PdfReader(str(_safe(rel)))
    text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
    return text[:limit]


def read_docx(rel: str, limit: int = 20000) -> str:
    try:
        import docx
    except ImportError:
        raise RuntimeError("Falta la librería python-docx (vuelve a construir la imagen: ./scripts/edi.sh rebuild)")
    d = docx.Document(str(_safe(rel)))
    return "\n".join(p.text for p in d.paragraphs)[:limit]


def read_any(rel: str) -> str:
    ext = Path(rel).suffix.lower()
    if ext == ".pdf":
        return read_pdf(rel)
    if ext == ".docx":
        return read_docx(rel)
    return read_text(rel)


def write_docx(rel_new: str, text: str):
    import docx
    d = docx.Document()
    for line in text.split("\n"):
        d.add_paragraph(line)
    d.save(str(_safe(rel_new)))


def _sibling_edi_name(rel: str) -> str:
    p = Path(rel)
    return str(p.with_name(f"{p.stem}_EDI{p.suffix}"))


def summarize(rel: str, session_id: int | None, pregunta: str = "") -> tuple[str | None, str | None]:
    """Resume (o responde algo sobre) un archivo de texto/PDF/Word. Regresa (respuesta, proveedor)."""
    try:
        text = read_any(rel)
    except NotAllowed as e:
        return str(e), None
    except Exception as e:
        return f"No pude abrir {rel}: {e}", None
    if not text.strip():
        return f"El archivo {rel} está vacío o no pude leerle texto (puede ser un escaneo sin OCR).", None
    ask = pregunta or "Resume de qué trata en 3 o 4 frases y dime si parece que piden alguna acción o entrega."
    out, prov = providers.chat(f"{ask}\n\nContenido del archivo «{Path(rel).name}»:\n{text[:8000]}",
                                max_tokens=280)
    return out, prov


def propose_edit(rel: str, instruction: str, session_id: int | None) -> dict:
    """Pide a un proveedor (de preferencia uno marcado 'code') que reescriba el archivo según la instrucción.
    No escribe nada todavía: guarda un borrador pendiente de confirmar, igual que un WhatsApp o correo."""
    try:
        current = read_text(rel) if Path(rel).suffix.lower() != ".docx" else read_docx(rel)
    except NotAllowed as e:
        return {"reply": str(e), "data": {}, "intent": "file_edit"}
    except Exception as e:
        return {"reply": f"No pude abrir {rel}: {e}", "data": {}, "intent": "file_edit"}
    ext = Path(rel).suffix.lower()
    tag = "code" if ext in CODE_EXT else "chat"
    prompt = (f"Este es el contenido actual de «{Path(rel).name}»:\n{current[:8000]}\n\n"
              f"Instrucción del usuario: {instruction}\n\n"
              "Responde ÚNICAMENTE con el contenido completo del archivo ya corregido/actualizado, "
              "sin explicaciones antes ni después, sin comillas ni ``` alrededor.")
    out, prov = providers.chat(prompt, max_tokens=1800, tag=tag,
                                system="Eres un asistente de programación/edición de documentos preciso y conservador. "
                                       "Cambias sólo lo que se te pide y mantienes el resto igual.")
    if not out:
        return {"reply": "No logré conectar con ningún proveedor de IA para hacer el cambio.", "data": {}, "intent": "file_edit"}
    preview = "\n".join(difflib.unified_diff(current.splitlines(), out.splitlines(),
                                             fromfile=rel, tofile=_sibling_edi_name(rel), lineterm=""))[:3000]
    eid = db.add_file_edit(session_id or 0, rel, "write", instruction, preview or "(sin cambios detectados)", out)
    lines = len(preview.splitlines())
    return {"reply": f"Ya tengo el cambio listo para {Path(rel).name} ({lines} líneas de diferencia). "
                     f"¿Lo guardo como {Path(_sibling_edi_name(rel)).name}? Dime «sí» o «cancela».",
            "data": {"file_edit_id": eid, "diff": preview}, "intent": "file_edit", "provider": prov}


def apply_edit(eid: int) -> dict:
    e = db.get_file_edit(eid)
    if not e or e["status"] != "pending":
        return {"reply": "Ya no tengo ese cambio pendiente.", "data": {}, "intent": "file_edit"}
    rel = e["path"]
    new_rel = _sibling_edi_name(rel)
    try:
        if Path(rel).suffix.lower() == ".docx":
            write_docx(new_rel, e["new_content"])
        else:
            _safe(new_rel).write_text(e["new_content"], encoding="utf-8")
    except Exception as ex:
        return {"reply": f"No pude guardar el archivo: {ex}", "data": {}, "intent": "file_edit"}
    db.set_file_edit_status(eid, "applied")
    return {"reply": f"Listo, guardé el cambio en {Path(new_rel).name}, junto al original.",
            "data": {"path": new_rel}, "intent": "file_edit"}


def discard_edit(eid: int) -> dict:
    db.set_file_edit_status(eid, "discarded")
    return {"reply": "Va, no hago el cambio.", "data": {}, "intent": "file_edit"}
